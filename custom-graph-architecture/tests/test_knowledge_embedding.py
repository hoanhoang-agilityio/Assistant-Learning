"""Tests for the knowledge base's embedding pipeline: chunking, ids, embedding."""

from pathlib import Path

import pytest
from docx import Document

from src.configs.config import settings
from src.models import KnowledgeChunk
from src.services import knowledge

SECTIONS = [
    ("Protein", "Aim for 1.6 to 2.2 g per kg of bodyweight per day."),
    ("Hydration", "Drink to thirst, and add electrolytes on long sessions."),
]


def _write(
    path: Path,
    *,
    title: str | None = None,
    preamble: str | None = None,
    sections: list[tuple[str, str]] | None = None,
) -> Path:
    """One knowledge document on disk, styled the way the shipped ones are."""
    document = Document()
    if title is not None:
        document.add_heading(title, level=1)
    if preamble is not None:
        document.add_paragraph(preamble)

    for heading, body in sections if sections is not None else SECTIONS:
        document.add_heading(heading, level=2)
        if body:
            document.add_paragraph(body)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


@pytest.fixture
def document(tmp_path: Path) -> Path:
    """A document with a title, prose belonging to no section, and two sections."""
    return _write(
        tmp_path / "Sports_Nutrition.docx",
        title="Sports Nutrition",
        preamble="Intro prose that belongs to no section.",
    )


def test_section_is_the_chunk(document: Path) -> None:
    """A `Heading 2` delimits a passage: the retrieval unit a question maps onto."""
    chunks = knowledge.chunk_document(document)

    assert [chunk.heading for chunk in chunks] == ["Protein", "Hydration"]


def test_preamble_is_dropped(document: Path) -> None:
    """Prose under no section has no heading to embed with it, so it is not a passage."""
    chunks = knowledge.chunk_document(document)

    assert not any("no section" in chunk.text for chunk in chunks)


def test_a_heading_with_no_body_is_not_a_passage(tmp_path: Path) -> None:
    """It is a container; embedded alone it matches questions and tells the reader nothing."""
    path = _write(
        tmp_path / "Training.docx",
        title="Training",
        sections=[("Recovery", ""), ("Volume", "Ten to twenty hard sets a week.")],
    )

    assert [chunk.heading for chunk in knowledge.chunk_document(path)] == ["Volume"]


def test_heading_is_embedded_with_the_body(document: Path) -> None:
    """Without its heading, "1.6 to 2.2 g per kg" is a number about nothing."""
    protein = knowledge.chunk_document(document)[0]

    assert protein.text.startswith("Protein\n\n")
    assert "1.6 to 2.2 g per kg" in protein.text


def test_title_becomes_the_citation(document: Path) -> None:
    """`source` is what the QA answer cites, so it comes from the document's own title."""
    chunks = knowledge.chunk_document(document)

    assert {chunk.source for chunk in chunks} == {"Sports Nutrition"}


def test_a_title_carrying_its_file_name_is_tidied(tmp_path: Path) -> None:
    """Two of the three shipped documents are titled `Fitness_Nutrition_Knowledge_Base`."""
    path = _write(
        tmp_path / "Fitness_Nutrition.docx", title="Fitness_Nutrition_Knowledge_Base"
    )

    assert (
        knowledge.chunk_document(path)[0].source == "Fitness Nutrition Knowledge Base"
    )


def test_source_falls_back_to_the_file_name(tmp_path: Path) -> None:
    """A document with no title still has to be citable."""
    path = _write(tmp_path / "injury_knowledge_base.docx")

    assert knowledge.chunk_document(path)[0].source == "injury knowledge base"


def test_ids_are_stable_across_runs(document: Path) -> None:
    """Re-seeding an unchanged document must update in place, not append a second copy."""
    first = knowledge.chunk_document(document)
    second = knowledge.chunk_document(document)

    assert [chunk.id for chunk in first] == [chunk.id for chunk in second]
    assert first[0].id == "sports-nutrition#protein"


def test_repeated_heading_gets_its_own_id(tmp_path: Path) -> None:
    """Two sections sharing a heading would otherwise overwrite each other on upsert."""
    path = _write(
        tmp_path / "Training.docx",
        sections=[("Volume", "Ten sets."), ("Volume", "Twelve sets.")],
    )

    assert [chunk.id for chunk in knowledge.chunk_document(path)] == [
        "training#volume",
        "training#volume-2",
    ]


def test_long_section_is_split_with_overlap(tmp_path: Path) -> None:
    """A section past the embedding budget is split; every part keeps the heading."""
    body = " ".join(
        f"Rule {index} covers one aspect of weekly training volume."
        for index in range(120)
    )
    path = _write(tmp_path / "Training.docx", sections=[("Volume", body)])

    chunks = knowledge.chunk_document(path)

    assert len(chunks) > 1
    assert all(chunk.text.startswith("Volume\n\n") for chunk in chunks)
    assert all(
        len(chunk.text) <= knowledge.MAX_CHUNK_CHARS + len("Volume\n\n")
        for chunk in chunks
    )
    assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))


def test_content_hash_tracks_the_embedded_text(document: Path, tmp_path: Path) -> None:
    """The hash is what lets the seeder skip passages it has already paid to embed."""
    unchanged = knowledge.chunk_document(document)[0]

    edited = _write(
        tmp_path / "edited" / "Sports_Nutrition.docx",
        title="Sports Nutrition",
        sections=[("Protein", "Aim for 1.8 to 2.4 g per kg of bodyweight per day.")],
    )
    changed = knowledge.chunk_document(edited)[0]

    assert changed.id == unchanged.id
    assert changed.content_hash != unchanged.content_hash


def test_chunk_directory_ignores_what_is_not_a_document(tmp_path: Path) -> None:
    """Word writes a lock file beside an open document, and the folder may hold notes."""
    _write(tmp_path / "Nutrition.docx", sections=[("Protein", "Eat it.")])
    _write(tmp_path / "~$Nutrition.docx", sections=[("Protein", "Not a document.")])
    (tmp_path / "notes.txt").write_text("Protein\n\nNot a document.\n")

    assert len(knowledge.chunk_directory(tmp_path)) == 1


def test_the_shipped_documents_chunk_into_citable_passages() -> None:
    """The pipeline is only worth its cost against the corpus actually seeded."""
    chunks = knowledge.chunk_directory()

    assert len(chunks) > 1
    assert len({chunk.id for chunk in chunks}) == len(chunks)
    assert all(chunk.source and chunk.heading and chunk.text for chunk in chunks)
    assert all("_" not in chunk.source for chunk in chunks)


async def test_embed_chunks_returns_table_shaped_rows(
    document: Path, monkeypatch
) -> None:
    """The pipeline's output is a row per passage, ready for the seeder to upsert."""

    class FakeEmbedder:
        async def aembed_documents(self, texts: list[str]) -> list[list[float]]:
            return [[0.1] * settings.KNOWLEDGE_EMBEDDING_DIM for _ in texts]

    monkeypatch.setattr(knowledge, "embedder", FakeEmbedder)
    rows = await knowledge.embed_chunks(knowledge.chunk_document(document))

    assert len(rows) == 2
    assert set(rows[0]) == set(KnowledgeChunk.model_fields)
    assert len(rows[0]["embedding"]) == settings.KNOWLEDGE_EMBEDDING_DIM


async def test_embedding_is_skipped_when_there_is_nothing_to_embed() -> None:
    """An empty batch must not cost an API call."""
    assert await knowledge.embed_chunks([]) == []
